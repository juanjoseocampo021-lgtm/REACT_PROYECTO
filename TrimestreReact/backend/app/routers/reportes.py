from datetime import datetime
from pathlib import Path
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from io import BytesIO

from ..database import get_db
from ..dependencies import obtener_usuario_actual, requerir_roles
from ..models import Venta, Factura, DetalleVenta

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])

BRAND_HEX = '#1a5c59'
BRAND_LIGHT_HEX = '#e6f4f1'
GOLD_HEX = '#d4a853'

_LOGO_PATH = Path(__file__).resolve().parents[1] / 'horizonte-logo.svg'
_logo_drawing = None

def _load_logo():
    global _logo_drawing
    if _logo_drawing is not None:
        return _logo_drawing
    try:
        from svglib.svglib import svg2rlg
        if _LOGO_PATH.exists():
            d = svg2rlg(str(_LOGO_PATH))
            if d:
                target_h = 40
                scale = target_h / d.height
                d.width = d.width * scale
                d.height = target_h
                d.scale(scale, scale)
                _logo_drawing = d
                return _logo_drawing
    except Exception:
        pass
    return None


def _crear_header_pdf(elements, styles, titulo, subtitulo=None):
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import Table, TableStyle, Spacer, Paragraph
    from reportlab.graphics.shapes import Drawing, Rect, String, Line

    d = Drawing(500, 70)
    d.add(Rect(0, 50, 500, 20, fillColor=colors.HexColor(BRAND_HEX), strokeColor=None))
    d.add(Rect(0, 47, 500, 3, fillColor=colors.HexColor(GOLD_HEX), strokeColor=None))

    logo = _load_logo()
    if logo:
        d.add(logo)
        text_x = 160
    else:
        text_x = 20

    d.add(String(text_x, 55, 'HORIZONTE VIAJES', fillColor=colors.white, fontSize=14, fontName='Helvetica-Bold'))
    d.add(String(350, 55, 'Documento oficial', fillColor=colors.white, fontSize=8, fontName='Helvetica'))
    d.add(Line(0, 44, 500, 44, strokeColor=colors.HexColor(BRAND_HEX), strokeWidth=0.5))
    elements.append(d)
    elements.append(Spacer(1, 8))

    title_style = styles['Title'].clone('ReportTitle')
    title_style.textColor = colors.HexColor(BRAND_HEX)
    title_style.fontSize = 16
    title_style.spaceAfter = 4
    elements.append(Paragraph(titulo, title_style))

    if subtitulo:
        sub_style = styles['Normal'].clone('ReportSub')
        sub_style.textColor = colors.HexColor('#4a5568')
        sub_style.fontSize = 10
        elements.append(Paragraph(subtitulo, sub_style))

    elements.append(Spacer(1, 6))


def _crear_footer_pdf(elements, styles):
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import Spacer
    from reportlab.graphics.shapes import Drawing, Rect, String, Line

    elements.append(Spacer(1, 20))
    d = Drawing(500, 30)
    d.add(Line(0, 20, 500, 20, strokeColor=colors.HexColor('#e2e8f0'), strokeWidth=0.5))
    d.add(String(0, 5, f'Horizonte Viajes — Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}',
                 fillColor=colors.HexColor('#718096'), fontSize=7, fontName='Helvetica'))
    d.add(String(350, 5, 'www.horizonteviajes.com', fillColor=colors.HexColor(BRAND_HEX), fontSize=7, fontName='Helvetica-Bold'))
    elements.append(d)


def _estilo_tabla_reportlab(table, brand, brand_light):
    from reportlab.lib import colors
    table.setStyle([
        ('BACKGROUND', (0, 0), (-1, 0), brand),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('FONTSIZE', (0, 1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 10),
        ('TOPPADDING', (0, 0), (-1, 0), 10),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 6),
        ('TOPPADDING', (0, 1), (-1, -1), 6),
        ('BACKGROUND', (0, 1), (-1, -1), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, brand_light]),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), brand_light),
        ('LINEABOVE', (0, -1), (-1, -1), 1.5, brand),
    ])


def _estilo_tabla_excel(ws, headers, data_rows, total_row_data, total_row_num, start_row=4):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    header_fill = PatternFill(start_color=BRAND_HEX.replace('#', ''), end_color=BRAND_HEX.replace('#', ''), fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11, name='Calibri')
    data_font = Font(size=10, name='Calibri')
    bold_font = Font(bold=True, size=10, name='Calibri')
    total_font = Font(bold=True, size=11, name='Calibri', color=BRAND_HEX.replace('#', ''))
    thin_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='thin', color='D0D0D0'),
        bottom=Side(style='thin', color='D0D0D0'),
    )
    gold_border = Border(
        left=Side(style='thin', color='D0D0D0'),
        right=Side(style='thin', color='D0D0D0'),
        top=Side(style='medium', color=BRAND_HEX.replace('#', '')),
        bottom=Side(style='double', color=BRAND_HEX.replace('#', '')),
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    alt_fill = PatternFill(start_color=BRAND_LIGHT_HEX.replace('#', ''), end_color=BRAND_LIGHT_HEX.replace('#', ''), fill_type="solid")

    for row_idx, data_row in enumerate(data_rows):
        excel_row = start_row + 1 + row_idx
        for col, value in enumerate(data_row, 1):
            cell = ws.cell(row=excel_row, column=col, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.font = data_font
            if row_idx % 2 == 1:
                cell.fill = alt_fill

    if total_row_data:
        for col, value in enumerate(total_row_data, 1):
            cell = ws.cell(row=total_row_num, column=col, value=value)
            cell.font = total_font
            cell.border = gold_border
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = PatternFill(start_color=BRAND_LIGHT_HEX.replace('#', ''), end_color=BRAND_LIGHT_HEX.replace('#', ''), fill_type="solid")


# ============================================================
# REPORTE DE VENTAS - PDF
# ============================================================
@router.get("/ventas/pdf")
def reporte_ventas_pdf(
    fecha: str = None,
    usuario: dict = Depends(requerir_roles("administrador", "empleado")),
    db: Session = Depends(get_db),
):
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        raise HTTPException(status_code=500, detail={"mensaje": "Librería reportlab no instalada"})

    if fecha:
        try:
            fecha_obj = datetime.strptime(fecha, "%Y-%m-%d")
            fecha_fin = fecha_obj.replace(hour=23, minute=59, second=59)
        except ValueError:
            fecha_obj = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
            fecha_fin = datetime.now()
    else:
        fecha_obj = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
        fecha_fin = datetime.now()

    ventas = (
        db.query(Venta)
        .filter(Venta.creado_en >= fecha_obj, Venta.creado_en <= fecha_fin)
        .order_by(Venta.creado_en.desc())
        .all()
    )

    BRAND = colors.HexColor(BRAND_HEX)
    BRAND_LIGHT = colors.HexColor(BRAND_LIGHT_HEX)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    elements = []

    _crear_header_pdf(elements, styles,
                      "Reporte Diario de Ventas",
                      f"<b>Fecha:</b> {fecha_obj.strftime('%d/%m/%Y')} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Generado:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    data = [['#', 'Cliente', 'Documento', 'Subtotal', 'IVA', 'Total', 'Estado', 'Hora']]
    for v in ventas:
        data.append([
            str(v.id),
            v.cliente_nombre or '',
            v.cliente_documento or '',
            f"${float(v.subtotal or 0):,.2f}",
            f"${float(v.impuestos or 0):,.2f}",
            f"${float(v.total or 0):,.2f}",
            (v.estado or '').capitalize(),
            v.creado_en.strftime('%H:%M') if v.creado_en else '',
        ])

    if len(data) > 1:
        total_sub = sum(float(v.subtotal or 0) for v in ventas)
        total_iva = sum(float(v.impuestos or 0) for v in ventas)
        total_gen = sum(float(v.total or 0) for v in ventas)
        data.append(['', '', '', '', '', '', '', ''])
        data.append(['', '', 'TOTAL', f'${total_sub:,.2f}', f'${total_iva:,.2f}', f'${total_gen:,.2f}', '', ''])

    table = Table(data, colWidths=[0.4*inch, 1.3*inch, 0.9*inch, 0.8*inch, 0.7*inch, 0.8*inch, 0.7*inch, 0.5*inch])
    _estilo_tabla_reportlab(table, BRAND, BRAND_LIGHT)
    elements.append(table)

    elements.append(Spacer(1, 10))
    resumen_style = styles['Normal'].clone('Resumen')
    resumen_style.textColor = colors.HexColor('#4a5568')
    resumen_style.fontSize = 9
    total_gen = sum(float(v.total or 0) for v in ventas)
    pagadas = sum(1 for v in ventas if v.estado == 'completada')
    pendientes = sum(1 for v in ventas if v.estado == 'pendiente')
    elements.append(Paragraph(f"<b>Resumen:</b> {len(ventas)} ventas &nbsp;|&nbsp; Completadas: {pagadas} &nbsp;|&nbsp; Pendientes: {pendientes} &nbsp;|&nbsp; Ingresos: <b>${total_gen:,.2f}</b>", resumen_style))

    _crear_footer_pdf(elements, styles)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=reporte_ventas_{fecha_obj.strftime('%Y%m%d')}.pdf"}
    )


# ============================================================
# REPORTE DE VENTAS - EXCEL
# ============================================================
@router.get("/ventas/excel")
def reporte_ventas_excel(
    fecha: str = None,
    usuario: dict = Depends(requerir_roles("administrador", "empleado")),
    db: Session = Depends(get_db),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail={"mensaje": "Librería openpyxl no instalada"})

    if fecha:
        try:
            fecha_obj = datetime.strptime(fecha, "%Y-%m-%d")
            fecha_fin = fecha_obj.replace(hour=23, minute=59, second=59)
        except ValueError:
            fecha_obj = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
            fecha_fin = datetime.now()
    else:
        fecha_obj = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
        fecha_fin = datetime.now()

    ventas = (
        db.query(Venta)
        .filter(Venta.creado_en >= fecha_obj, Venta.creado_en <= fecha_fin)
        .order_by(Venta.creado_en.desc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"

    brand_fill = PatternFill(start_color=BRAND_HEX.replace('#', ''), end_color=BRAND_HEX.replace('#', ''), fill_type="solid")
    gold_fill = PatternFill(start_color=GOLD_HEX.replace('#', ''), end_color=GOLD_HEX.replace('#', ''), fill_type="solid")
    brand_font_lg = Font(bold=True, size=16, color=BRAND_HEX.replace('#', ''), name='Calibri')
    brand_font_md = Font(bold=True, size=12, color=BRAND_HEX.replace('#', ''), name='Calibri')
    subtitle_font = Font(italic=True, size=10, color='666666', name='Calibri')

    ws.merge_cells('A1:K1')
    c = ws.cell(row=1, column=1, value='HORIZONTE VIAJES')
    c.font = brand_font_lg
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:K2')
    c2 = ws.cell(row=2, column=1, value='Reporte de Ventas')
    c2.font = brand_font_md
    c2.alignment = Alignment(horizontal='left')

    ws.merge_cells('A3:K3')
    c3 = ws.cell(row=3, column=1, value=f'Fecha: {fecha_obj.strftime("%d/%m/%Y")}  |  Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}')
    c3.font = subtitle_font

    ws.row_dimensions[4].height = 5
    gold_row = ws.cell(row=4, column=1)
    gold_row.fill = gold_fill

    headers = ['ID', 'Cliente', 'Documento', 'Correo', 'Teléfono', 'Subtotal', 'Impuestos', 'Descuento', 'Total', 'Estado', 'Fecha/Hora']
    data_rows = []
    for v in ventas:
        data_rows.append([
            v.id, v.cliente_nombre or '', v.cliente_documento or '',
            v.cliente_correo or '', v.cliente_telefono or '',
            float(v.subtotal or 0), float(v.impuestos or 0), float(v.descuento or 0),
            float(v.total or 0), (v.estado or '').capitalize(),
            v.creado_en.strftime('%d/%m/%Y %H:%M') if v.creado_en else ''
        ])

    total_gen = sum(float(v.total or 0) for v in ventas)
    total_sub = sum(float(v.subtotal or 0) for v in ventas)
    total_iva = sum(float(v.impuestos or 0) for v in ventas)
    total_desc = sum(float(v.descuento or 0) for v in ventas)
    total_row = [''] * 4 + ['TOTALES', total_sub, total_iva, total_desc, total_gen, '', '']
    total_row_num = 5 + len(data_rows) + 1

    _estilo_tabla_excel(ws, headers, data_rows, total_row, total_row_num, start_row=5)

    ws.row_dimensions[total_row_num + 1].height = 5
    ws.cell(row=total_row_num + 1, column=1).fill = gold_fill

    ws.merge_cells(f'A{total_row_num + 2}:K{total_row_num + 2}')
    ws.cell(row=total_row_num + 2, column=1, value=f'Horizonte Viajes — {len(ventas)} registros — Ingresos totales: ${total_gen:,.2f}').font = Font(italic=True, size=9, color='999999', name='Calibri')

    for col_idx in range(1, len(headers) + 1):
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(5, total_row_num))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 28)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=reporte_ventas_{fecha_obj.strftime('%Y%m%d')}.xlsx"}
    )


# ============================================================
# FACTURA INDIVIDUAL - PDF
# ============================================================
@router.get("/facturas/pdf/{factura_id}")
def factura_individual_pdf(
    factura_id: int,
    usuario: dict = Depends(requerir_roles("administrador", "empleado")),
    db: Session = Depends(get_db),
):
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, HRFlowable
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        raise HTTPException(status_code=500, detail={"mensaje": "Librería reportlab no instalada"})

    factura = db.query(Factura).filter(Factura.id == factura_id).first()
    if not factura:
        raise HTTPException(status_code=404, detail={"mensaje": "Factura no encontrada"})

    venta = db.query(Venta).filter(Venta.id == factura.venta_id).first()
    detalles = db.query(DetalleVenta).filter(DetalleVenta.venta_id == factura.venta_id).all() if venta else []

    BRAND = colors.HexColor(BRAND_HEX)
    BRAND_LIGHT = colors.HexColor(BRAND_LIGHT_HEX)
    GOLD = colors.HexColor(GOLD_HEX)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.4*inch, bottomMargin=0.4*inch, leftMargin=0.6*inch, rightMargin=0.6*inch)
    styles = getSampleStyleSheet()
    elements = []

    _crear_header_pdf(elements, styles,
                      f"Factura {factura.numero_factura}",
                      f"<b>Fecha:</b> {factura.creado_en.strftime('%d/%m/%Y %H:%M') if factura.creado_en else 'N/A'} &nbsp;&nbsp;|&nbsp;&nbsp; <b>Estado:</b> {(factura.estado or '').upper()}")

    info_data = [
        [Paragraph(f"<b>Factura N°:</b> {factura.numero_factura}", styles['Normal']),
         Paragraph(f"<b>Venta ID:</b> #{factura.venta_id}", styles['Normal'])],
    ]
    info_table = Table(info_data, colWidths=[3.5*inch, 3.5*inch])
    info_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(info_table)
    elements.append(Spacer(1, 12))

    elements.append(Paragraph("<b>Datos del Cliente</b>", styles['Heading3']))
    client_data = [
        ['Nombre', factura.cliente_nombre or 'N/A'],
        ['Documento', factura.cliente_documento or 'N/A'],
        ['Correo', factura.cliente_correo or 'N/A'],
    ]
    if venta:
        client_data.append(['Teléfono', venta.cliente_telefono or 'N/A'])
    client_table = Table(client_data, colWidths=[1.5*inch, 5.5*inch])
    client_table.setStyle(TableStyle([
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, 0), (0, -1), BRAND_LIGHT),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(client_table)
    elements.append(Spacer(1, 15))

    elements.append(Paragraph("<b>Detalle de productos/servicios</b>", styles['Heading3']))
    detail_data = [['Cant.', 'Producto/Servicio', 'Precio Unit.', 'Subtotal']]
    for det in detalles:
        detail_data.append([
            str(det.cantidad or 1),
            det.nombre_item or 'N/A',
            f"${float(det.precio_unitario or 0):,.2f}",
            f"${float(det.subtotal or 0):,.2f}",
        ])
    if not detalles:
        detail_data.append(['', 'Sin detalles disponibles', '', ''])

    detail_table = Table(detail_data, colWidths=[0.8*inch, 3.2*inch, 1.5*inch, 1.5*inch])
    detail_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), BRAND),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 10),
        ('ALIGN', (-1, 1), (-1, -1), 'RIGHT'),
        ('ALIGN', (-2, 1), (-2, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, BRAND_LIGHT]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
    ]))
    elements.append(detail_table)
    elements.append(Spacer(1, 15))

    totals_data = [
        ['Subtotal:', f"${float(factura.subtotal or 0):,.2f}"],
        ['Impuestos (IVA 19%):', f"${float(factura.impuestos or 0):,.2f}"],
        ['Descuento:', f"-${float(factura.descuento or 0):,.2f}"],
        ['TOTAL:', f"${float(factura.total or 0):,.2f}"],
    ]
    totals_table = Table(totals_data, colWidths=[5*inch, 2*inch])
    totals_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (0, -1), 'RIGHT'),
        ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, -1), (-1, -1), 12),
        ('TEXTCOLOR', (0, -1), (-1, -1), BRAND),
        ('LINEABOVE', (0, -1), (-1, -1), 2, GOLD),
        ('TOPPADDING', (0, -1), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(totals_table)

    _crear_footer_pdf(elements, styles)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=factura_{factura.numero_factura}.pdf"}
    )


# ============================================================
# REPORTE DE FACTURAS - PDF
# ============================================================
@router.get("/facturas/pdf")
def reporte_facturas_pdf(
    usuario: dict = Depends(requerir_roles("administrador", "empleado")),
    db: Session = Depends(get_db),
):
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet
    except ImportError:
        raise HTTPException(status_code=500, detail={"mensaje": "Librería reportlab no instalada"})

    facturas = db.query(Factura).order_by(Factura.creado_en.desc()).all()

    BRAND = colors.HexColor(BRAND_HEX)
    BRAND_LIGHT = colors.HexColor(BRAND_LIGHT_HEX)

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5*inch, bottomMargin=0.5*inch)
    styles = getSampleStyleSheet()
    elements = []

    _crear_header_pdf(elements, styles,
                      "Reporte General de Facturas",
                      f"<b>Generado:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}")

    data = [['N° Factura', 'Cliente', 'Documento', 'Subtotal', 'IVA', 'Total', 'Estado', 'Fecha']]
    for f in facturas:
        data.append([
            f.numero_factura or '',
            f.cliente_nombre or '',
            f.cliente_documento or '',
            f"${float(f.subtotal or 0):,.2f}",
            f"${float(f.impuestos or 0):,.2f}",
            f"${float(f.total or 0):,.2f}",
            (f.estado or '').capitalize(),
            f.creado_en.strftime('%d/%m/%Y') if f.creado_en else '',
        ])

    if len(data) > 1:
        total_gen = sum(float(f.total or 0) for f in facturas)
        data.append(['', '', '', '', '', '', '', ''])
        data.append(['', '', 'TOTAL', '', '', f'${total_gen:,.2f}', '', ''])

    table = Table(data, colWidths=[0.9*inch, 1.3*inch, 0.8*inch, 0.8*inch, 0.7*inch, 0.8*inch, 0.7*inch, 0.8*inch])
    _estilo_tabla_reportlab(table, BRAND, BRAND_LIGHT)
    elements.append(table)

    elements.append(Spacer(1, 10))
    total_fact = sum(float(f.total or 0) for f in facturas)
    pagadas = sum(1 for f in facturas if f.estado == 'pagada')
    pendientes = sum(1 for f in facturas if f.estado == 'pendiente')
    anuladas = sum(1 for f in facturas if f.estado == 'anulada')
    resumen_style = styles['Normal'].clone('Resumen2')
    resumen_style.textColor = colors.HexColor('#4a5568')
    resumen_style.fontSize = 9
    elements.append(Paragraph(f"<b>Resumen:</b> {len(facturas)} facturas &nbsp;|&nbsp; Pagadas: {pagadas} &nbsp;|&nbsp; Pendientes: {pendientes} &nbsp;|&nbsp; Anuladas: {anuladas} &nbsp;|&nbsp; Total: <b>${total_fact:,.2f}</b>", resumen_style))

    _crear_footer_pdf(elements, styles)

    doc.build(elements)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=reporte_facturas.pdf"}
    )


# ============================================================
# REPORTE DE FACTURAS - EXCEL
# ============================================================
@router.get("/facturas/excel")
def reporte_facturas_excel(
    usuario: dict = Depends(requerir_roles("administrador", "empleado")),
    db: Session = Depends(get_db),
):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise HTTPException(status_code=500, detail={"mensaje": "Librería openpyxl no instalada"})

    facturas = db.query(Factura).order_by(Factura.creado_en.desc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Facturas"

    brand_fill = PatternFill(start_color=BRAND_HEX.replace('#', ''), end_color=BRAND_HEX.replace('#', ''), fill_type="solid")
    gold_fill = PatternFill(start_color=GOLD_HEX.replace('#', ''), end_color=GOLD_HEX.replace('#', ''), fill_type="solid")
    brand_font_lg = Font(bold=True, size=16, color=BRAND_HEX.replace('#', ''), name='Calibri')
    brand_font_md = Font(bold=True, size=12, color=BRAND_HEX.replace('#', ''), name='Calibri')
    subtitle_font = Font(italic=True, size=10, color='666666', name='Calibri')

    ws.merge_cells('A1:J1')
    c = ws.cell(row=1, column=1, value='HORIZONTE VIAJES')
    c.font = brand_font_lg
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:J2')
    c2 = ws.cell(row=2, column=1, value='Reporte General de Facturas')
    c2.font = brand_font_md

    ws.merge_cells('A3:J3')
    c3 = ws.cell(row=3, column=1, value=f'Generado: {datetime.now().strftime("%d/%m/%Y %H:%M")}')
    c3.font = subtitle_font

    ws.row_dimensions[4].height = 5
    ws.cell(row=4, column=1).fill = gold_fill

    headers = ['N° Factura', 'Cliente', 'Documento', 'Correo', 'Subtotal', 'Impuestos', 'Descuento', 'Total', 'Estado', 'Fecha']
    data_rows = []
    for f in facturas:
        data_rows.append([
            f.numero_factura or '', f.cliente_nombre or '', f.cliente_documento or '',
            f.cliente_correo or '',
            float(f.subtotal or 0), float(f.impuestos or 0), float(f.descuento or 0),
            float(f.total or 0), (f.estado or '').capitalize(),
            f.creado_en.strftime('%d/%m/%Y %H:%M') if f.creado_en else ''
        ])

    total_gen = sum(float(f.total or 0) for f in facturas)
    total_sub = sum(float(f.subtotal or 0) for f in facturas)
    total_iva = sum(float(f.impuestos or 0) for f in facturas)
    total_desc = sum(float(f.descuento or 0) for f in facturas)
    total_row = [''] * 3 + ['TOTALES', total_sub, total_iva, total_desc, total_gen, '', '']
    total_row_num = 5 + len(data_rows) + 1

    _estilo_tabla_excel(ws, headers, data_rows, total_row, total_row_num, start_row=5)

    ws.row_dimensions[total_row_num + 1].height = 5
    ws.cell(row=total_row_num + 1, column=1).fill = gold_fill

    ws.merge_cells(f'A{total_row_num + 2}:J{total_row_num + 2}')
    ws.cell(row=total_row_num + 2, column=1, value=f'Horizonte Viajes — {len(facturas)} facturas — Total: ${total_gen:,.2f}').font = Font(italic=True, size=9, color='999999', name='Calibri')

    for col_idx in range(1, len(headers) + 1):
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or '')) for r in range(5, total_row_num))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 28)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=reporte_facturas.xlsx"}
    )
